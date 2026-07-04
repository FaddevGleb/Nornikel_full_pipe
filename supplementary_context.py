"""Parse user-uploaded Excel documents into supplementary ACCELMAT context."""

from __future__ import annotations

from dataclasses import dataclass, field
from math import isnan
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from logging_utils import get_logger

logger = get_logger("supplementary_context")

DEFAULT_MAX_TOTAL_CHARS = 32_000
DEFAULT_MAX_CHARS_PER_SHEET = 8_000

SUPPLEMENTARY_HEADER = (
    "\n\n### Дополнительные документы (загруженные пользователем таблицы Excel):\n"
    "Считай их авторитетными операционными данными (содержания, анализы, схемы, затраты). "
    "Гипотезы должны согласовываться с этими цифрами, где это уместно.\n"
)


@dataclass
class ParsedSheet:
    name: str
    markdown: str
    char_count: int


@dataclass
class ParsedDocument:
    filename: str
    path: str
    sheets: list[ParsedSheet] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return sum(sheet.char_count for sheet in self.sheets)


@dataclass
class SupplementaryContextResult:
    documents: list[dict[str, Any]]
    formatted_text: str

    def to_dict(self) -> dict[str, Any]:
        return {
            "documents": self.documents,
            "formatted_text": self.formatted_text,
        }


def _escape_md_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ")


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and isnan(value):
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _build_merged_lookup(sheet: Worksheet) -> dict[tuple[int, int], tuple[int, int, str]]:
    lookup: dict[tuple[int, int], tuple[int, int, str]] = {}
    for merge_range in sheet.merged_cells.ranges:
        min_row, min_col = merge_range.min_row, merge_range.min_col
        anchor_value = _cell_to_str(sheet.cell(min_row, min_col).value)
        for row in range(merge_range.min_row, merge_range.max_row + 1):
            for col in range(merge_range.min_col, merge_range.max_col + 1):
                lookup[(row, col)] = (min_row, min_col, anchor_value)
    return lookup


def _sheet_dimensions(sheet: Worksheet) -> tuple[int, int]:
    return sheet.max_row or 0, sheet.max_column or 0


def _collect_grid(sheet: Worksheet, df: pd.DataFrame) -> list[list[str]]:
    max_row, max_col = _sheet_dimensions(sheet)
    if max_row == 0 or max_col == 0:
        return []

    merged = _build_merged_lookup(sheet)
    grid: list[list[str]] = []

    for row_idx in range(1, max_row + 1):
        row_values: list[str] = []
        for col_idx in range(1, max_col + 1):
            if (row_idx, col_idx) in merged:
                _, _, value = merged[(row_idx, col_idx)]
            else:
                df_row = row_idx - 1
                df_col = col_idx - 1
                if df_row < len(df.index) and df_col < len(df.columns):
                    value = _cell_to_str(df.iat[df_row, df_col])
                else:
                    value = _cell_to_str(sheet.cell(row_idx, col_idx).value)
            row_values.append(value)
        grid.append(row_values)

    return grid


def _trim_empty_edges(grid: list[list[str]]) -> list[list[str]]:
    if not grid:
        return grid

    def row_empty(row: list[str]) -> bool:
        return all(not cell for cell in row)

    while grid and row_empty(grid[-1]):
        grid.pop()
    while grid and all(not row[0] for row in grid if row):
        for row in grid:
            if row:
                row.pop(0)
    while grid and all(not row[-1] for row in grid if row):
        for row in grid:
            if row:
                row.pop()
    return grid


def _grid_to_markdown(grid: list[list[str]]) -> str:
    if not grid:
        return ""

    col_count = max(len(row) for row in grid)
    normalized = [row + [""] * (col_count - len(row)) for row in grid]

    header = normalized[0]
    body = normalized[1:] if len(normalized) > 1 else [[""] * col_count]

    lines = [
        "| " + " | ".join(_escape_md_cell(c) for c in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(_escape_md_cell(c) for c in row) + " |")
    return "\n".join(lines)


def _truncate_text(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    return text[: max(0, max_chars - len("[truncated]"))] + "[truncated]"


def parse_excel_workbook(path: str | Path) -> ParsedDocument:
    path = Path(path)
    workbook = load_workbook(path, data_only=True)
    sheets_data: dict[str, pd.DataFrame] = pd.read_excel(path, sheet_name=None, header=None)

    parsed_sheets: list[ParsedSheet] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        df = sheets_data.get(sheet_name, pd.DataFrame())
        grid = _trim_empty_edges(_collect_grid(sheet, df))
        if not grid:
            continue
        markdown = _grid_to_markdown(grid)
        parsed_sheets.append(
            ParsedSheet(name=sheet_name, markdown=markdown, char_count=len(markdown))
        )

    return ParsedDocument(
        filename=path.name,
        path=str(path),
        sheets=parsed_sheets,
    )


def format_context_section(text: str) -> str:
    stripped = text.strip()
    if not stripped:
        return ""
    return f"{SUPPLEMENTARY_HEADER}{stripped}\n"


def append_supplementary_to_prompt(prompt: str, supplementary_context: str | None) -> str:
    section = format_context_section(supplementary_context or "")
    if not section:
        return prompt
    return f"{prompt}{section}"


def append_supplementary_to_system(system_content: str, supplementary_context: str | None) -> str:
    section = format_context_section(supplementary_context or "")
    if not section:
        return system_content
    return f"{system_content}{section}"


def build_supplementary_context(
    paths: list[str | Path],
    *,
    base_dir: str | Path | None = None,
    max_total_chars: int = DEFAULT_MAX_TOTAL_CHARS,
    max_chars_per_sheet: int = DEFAULT_MAX_CHARS_PER_SHEET,
) -> SupplementaryContextResult:
    if not paths:
        return SupplementaryContextResult(documents=[], formatted_text="")

    base = Path(base_dir) if base_dir else None
    documents: list[ParsedDocument] = []
    body_parts: list[str] = []
    total_chars = 0
    budget_exhausted = False

    for raw_path in paths:
        if budget_exhausted:
            break

        path = Path(raw_path)
        if not path.is_absolute() and base is not None:
            path = base / path
        if not path.exists():
            logger.warning("[Supplementary] Skipping missing document: %s", path)
            continue
        if path.suffix.lower() != ".xlsx":
            logger.warning("[Supplementary] Skipping unsupported extension: %s", path)
            continue

        try:
            parsed_doc = parse_excel_workbook(path)
        except Exception as exc:
            logger.warning("[Supplementary] Failed to parse %s: %s", path, exc)
            continue

        if not parsed_doc.sheets:
            logger.warning("[Supplementary] No sheets with data in %s", path)
            continue

        documents.append(parsed_doc)
        body_parts.append(f"#### File: {parsed_doc.filename}\n")

        for sheet in parsed_doc.sheets:
            if budget_exhausted:
                break

            sheet_text = f"##### Sheet: {sheet.name}\n{sheet.markdown}"
            sheet_text = _truncate_text(sheet_text, max_chars_per_sheet)

            remaining = max_total_chars - total_chars
            if remaining <= 0:
                body_parts.append("[truncated: total character budget exceeded]\n")
                budget_exhausted = True
                break

            if len(sheet_text) > remaining:
                sheet_text = _truncate_text(sheet_text, remaining)
                budget_exhausted = True

            body_parts.append(sheet_text)
            body_parts.append("\n")
            total_chars += len(sheet_text)

    formatted_body = "".join(body_parts).strip()
    formatted_text = format_context_section(formatted_body) if formatted_body else ""

    doc_meta = [
        {
            "filename": doc.filename,
            "path": doc.path,
            "sheets": [sheet.name for sheet in doc.sheets],
            "char_count": doc.char_count,
        }
        for doc in documents
    ]

    logger.info(
        "[Supplementary] Built context from %d document(s), %d chars",
        len(documents),
        len(formatted_text),
    )
    return SupplementaryContextResult(documents=doc_meta, formatted_text=formatted_text)
